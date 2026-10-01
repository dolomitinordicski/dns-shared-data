"""Read immutable 2025-26 snapshots. Never modify or execute workbook macros."""
import json, pathlib, hashlib, sys
import openpyxl

ROOT = pathlib.Path(sys.argv[1])
sources = [
    ('orders', '01-ALL-TICKETS-2025-26.xlsx', '825d23a7b732ce3ed8b70e348f58cda97b3e192dbbb50078d2fc7556e552136a'),
    ('kp', '02-KP-Artificial-vs.-natural-snow-2025-26.xlsm', '0a01f3eaab87f38c952d637756747779e0a40583c150c5fb1b562dccb2f9499b'),
    ('sales', '03-VERKAUFSTATISTIK-STATISTICHE-DI-VENDITA-2025-26.xlsx', '61e19d8cdaafdb06a38cefa2f8b26340689d943eaf16c3692b2b144ac5b64504'),
]
records, archives = [], []
season = '2025-26'
partners = {
 'antholzertal': ('antholzertal','antholzertal'), 'biathlon-antholz': ('biathlon-antholz','antholzertal'),
 'gsiesertal-welsberg-taisten': ('gsiesertal-welsberg-taisten','gsiesertal-welsberg-taisten'),
 'tv-toblach': ('tv-toblach','drei-zinnen'), 'tv-sexten': ('tv-sexten','drei-zinnen'),
 'tv-innichen': ('tv-innichen','drei-zinnen'), 'tv-niederdorf': ('tv-niederdorf','drei-zinnen'),
 'tv-prags': ('tv-prags','drei-zinnen'), 'tvb-osttirol': ('tvb-osttirol','osttirol'),
 'val-comelico': ('val-comelico','val-comelico'), 'servizi-ampezzo': ('servizi-ampezzo','cortina-d-ampezzo'),
 'sand-in-taufers': ('sand-in-taufers','ahrntal'), 'ahrntal': ('ahrntal','ahrntal'),
 'val-gardena': ('val-gardena','seiser-alm-dolomites-val-gardena'),
 'seiser-alm-marketing': ('seiser-alm-marketing','seiser-alm-dolomites-val-gardena'),
 'dolomiti-nordicski': ('dolomiti-nordicski',''),
}

def add(domain, key, partner, sheet, facts, label=None):
 org, area = partners[partner]
 records.append(dict(id=f'{season}__{domain}__{key}', seasonId=season, domain=domain,
  organizationId=org, reportingAreaId=area, label=label or partner, sheet=sheet,
  facts=facts, readOnly=True, provenance=dict(sourceSystem='legacy-sheet', methodVersion=1,
  dataStatus='verified-with-notes', sourceRecordId=sheet)))

books = {}
for domain, filename, expected in sources:
 path=ROOT/filename
 digest=hashlib.sha256(path.read_bytes()).hexdigest()
 if digest != expected: raise ValueError(f'Source changed: {filename}')
 wb=openpyxl.load_workbook(path,data_only=True)
 formulas=openpyxl.load_workbook(path,data_only=False)
 books[domain]=wb
 for index, sheet in enumerate(wb):
  cells={}
  for row in sheet:
   for cell in row:
    original=formulas[sheet.title][cell.coordinate]
    if cell.value is None and original.value is None: continue
    cells[cell.coordinate]={'value':cell.value, 'formula':original.value if original.data_type=='f' else None}
  archives.append(dict(id=f'{season}__{domain}__{index}',seasonId=season,domain=domain,
   filename=filename,sha256=digest,sheet=sheet.title,cells=cells))

order_partners=['antholzertal','biathlon-antholz','gsiesertal-welsberg-taisten','tv-toblach',
 'tv-niederdorf','tv-innichen','tv-sexten','tv-prags','tvb-osttirol','val-comelico',
 'servizi-ampezzo','sand-in-taufers','ahrntal','val-gardena','seiser-alm-marketing']
for category,sheetname,start,end in [('wristband','Armbänder-braccialetti',3,17),
 ('ticket','Wochen- und Saisonkarten-settim',5,20)]:
 s=books['orders'][sheetname]
 for row in range(start,end+1):
  partner=(order_partners+['dolomiti-nordicski'])[row-start]
  facts=[dict(item=str(s.cell(2 if category=='wristband' else 3,c).value),
   quantity=s.cell(row,c).value,sourceCell=s.cell(row,c).coordinate) for c in range(2,10 if category=='wristband' else 9)]
  add('orders',category+'__'+partner,partner,sheetname,facts,str(s.cell(row,1).value))
  total=s.cell(row,10 if category=='wristband' else 9).value
  assert sum(f['quantity'] or 0 for f in facts)==total,(sheetname,row)

s=books['kp']['KP Artificial vs. natural ']
kp_partners=['antholzertal','biathlon-antholz','gsiesertal-welsberg-taisten','tv-toblach',
 'tv-sexten','tv-innichen','tv-niederdorf','tv-prags','tvb-osttirol','tvb-osttirol',
 'servizi-ampezzo','val-comelico','ahrntal','sand-in-taufers','seiser-alm-marketing','val-gardena']
for row,partner in enumerate(kp_partners,4):
 facts=[dict(date=date,referenceKm=s.cell(row,2).value,naturalKm=s.cell(row,c).value,
  artificialKm=s.cell(row,c+1).value,sourceCell=s.cell(row,c).coordinate)
  for date,c in [('2025-12-23',3),('2026-01-06',5),('2026-01-20',7)]]
 add('kp',str(row),partner,s.title,facts,str(s.cell(row,1).value))

# Leaf partner blocks only. Region totals are independent reported controls, never additive facts.
blocks=[('ANTHOLZERTAL','antholzertal',46),('ANTHOLZERTAL','biathlon-antholz',63),
 ('GSIESERTAL','gsiesertal-welsberg-taisten',24),('3 ZINNEN DOLOMITES','tv-toblach',46),
 ('3 ZINNEN DOLOMITES','tv-sexten',63),('3 ZINNEN DOLOMITES','tv-innichen',80),
 ('3 ZINNEN DOLOMITES','tv-niederdorf',97),('3 ZINNEN DOLOMITES','tv-prags',114),
 ('OSTTIROL','tvb-osttirol',24),('AHRNTAL','ahrntal',46),('AHRNTAL','sand-in-taufers',63),
 ('SEISER ALM DOLOMITES VAL GARDEN','seiser-alm-marketing',45),
 ('SEISER ALM DOLOMITES VAL GARDEN','val-gardena',62),('CORTINA','servizi-ampezzo',24),
 ('COMELICO','val-comelico',24)]
for sheetname,partner,start in blocks:
 s=books['sales'][sheetname]; facts=[]
 seasonal_offset=7 if start>=45 else 9
 for offset,product in [(0,'day'),(1,'wk-area'),(2,'wk-dns'),(seasonal_offset,'sk-area'),(seasonal_offset+1,'sk-dns'),(seasonal_offset+2,'sk-instructor')]:
  for c in [2,4,6]:
   row=start+offset; q=s.cell(row,c).value; amount=s.cell(row,c+1).value
   if q is None: continue # Missing is not zero.
   assert isinstance(q,(int,float)) and q>=0 and q==int(q)
   channel=['official','online','track'][(c-2)//2] if offset<3 else ('complimentary' if c==6 else 'official')
   period='regular' if offset<3 or c==4 else ('presale' if c==2 else 'unspecified')
   facts.append(dict(productCode=product,salesChannel=channel,salesPeriod=period,
    quantity=int(q),amount=amount,currency='EUR',sourceCell=s.cell(row,c).coordinate))
 add('sales',partner,partner,sheetname,facts,str(s.cell(start-4 if start>=45 else start-5,1).value))


for sheetname,partner,start in blocks:
 if partner in ['biathlon-antholz','tv-sexten','tv-innichen','tv-niederdorf','tv-prags','sand-in-taufers','val-gardena']: continue
 sheet=books['sales'][sheetname]; shift=1 if sheetname.startswith('SEISER') else 0
 facts=[]
 for column,rows in [(1,range(6-shift,16-shift)),(4,range(6-shift,11-shift))]:
  for row in rows:
   label=sheet.cell(row,column).value; price=sheet.cell(row,column+1).value
   if label is not None: facts.append(dict(item=str(label),unitPrice=price,currency='EUR',sourceCell=sheet.cell(row,column+1).coordinate))
 add('pricing',partner,partner,sheetname,facts)

controls=[]
analysis=books['sales']['DNS ANALYSE']
areas=['antholzertal','gsiesertal-welsberg-taisten','drei-zinnen','osttirol','ahrntal',
 'seiser-alm-dolomites-val-gardena','cortina-d-ampezzo','val-comelico']
for row,area in enumerate(areas,8):
 related=[r for r in records if r['domain']=='sales' and r['reportingAreaId']==area]
 quantity=sum(f['quantity'] for r in related for f in r['facts'])
 amount=sum(f['amount'] or 0 for r in related for f in r['facts'])
 controls.append(dict(reportingAreaId=area,reportedQuantity=analysis.cell(row,15).value,
  reportedAmount=analysis.cell(row,16).value,detailQuantity=quantity,detailAmount=amount,
  quantityDifference=analysis.cell(row,15).value-quantity,
  amountDifference=analysis.cell(row,16).value-amount,sourceCell=f'DNS ANALYSE!O{row}:P{row}'))
assert sum(c['quantityDifference'] for c in controls)==34
assert sum(c['amountDifference'] for c in controls)==2745
assert all(c['quantityDifference']==0 and c['amountDifference']==0 for c in controls[:5])
summary=dict(seasonId=season,readOnly=True,reportedQuantity=analysis['O16'].value,
 reportedAmount=analysis['P16'].value,controls=controls,
 notes=['Reported totals and leaf details are separate measures. Never sum both.',
 'Copied/template/previous-season sheets are archived, excluded from 2025-26 facts.',
 'KP analysed data has stale dates, inconsistent percentages and reference kilometres. Use dated source observations.',
 'KP cumulative milestones and unique reference kilometres must not be summed across dates.',
 'Osttirol and Obertilliach retained separately; overlap is not inferred.',
 'Sales source mentions Osttirol March online pending and Niederdorf commissions. Source preserved as supplied.'])
payload=dict(schemaVersion=1,seasonId=season,records=records,archives=archives,summary=summary)
pathlib.Path(sys.argv[2]).write_text(json.dumps(payload,ensure_ascii=False,default=str))
print(json.dumps({'records':len(records),'sourceSheets':len(archives),'controls':controls}))
